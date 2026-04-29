from urllib import request
from django.http import HttpResponseForbidden
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, logout
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.views.decorators.cache import never_cache
from .forms import SignUpForm, UserForm
from .models import *
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from django.contrib.auth.views import LoginView
from django.contrib.auth.forms import AuthenticationForm
from django.core.cache import cache
from django.http import JsonResponse


# Creating the helper function for report data
def get_report_data():
    users_count = User.objects.count()
    teams_count = Team.objects.count()
    departments_count = Department.objects.count()
    messages_count = Message.objects.count()
    conversations_count = Conversation.objects.count()

    recent_messages = Message.objects.select_related('sender', 'conversation').order_by('-timestamp')[:10]

    return {
        'users_count': users_count,
        'teams_count': teams_count,
        'departments_count': departments_count,
        'messages_count': messages_count,
        'conversations_count': conversations_count,
        'recent_messages': recent_messages,
    }

def admin_required_message(request):
    return render(request, 'admin_only.html')

# Making the chats grouped
def get_grouped_chat_data():
    conversations = Conversation.objects.prefetch_related('participants').order_by('id')
    grouped_data = []

    for convo in conversations:
        participants = list(convo.participants.all())
        participant_names = ", ".join([user.username for user in participants])

        messages = Message.objects.filter(conversation=convo).select_related('sender').order_by('timestamp')

        grouped_data.append({
            'conversation': convo,
            'participant_names': participant_names,
            'messages': messages,
        })

    return grouped_data


# Create views here.
def signup(request):
    if request.method == 'POST':
        form = SignUpForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            return redirect('/')
    else:
        form = SignUpForm()
    return render(request, 'registration/signup.html', {'form': form})

@login_required
@never_cache
def home(request):
    return render(request, 'home.html')

def user_logout(request):
    logout(request)
    return redirect('/accounts/login/')



# login view with failed password lockout

class LockedLoginView(LoginView):
    template_name = 'registration/login.html'
    authentication_form = AuthenticationForm

    def dispatch(self, request, *args, **kwargs):
        # check lock before password validation
        if request.method == 'POST':
            username = request.POST.get('username', '')

            # only lock real user accounts
            if User.objects.filter(username=username).exists():
                lock_key = f'login_lock_{username}'

                # block user if account is locked
                if cache.get(lock_key):
                    form = self.get_form()
                    form.add_error(None, 'Too many failed login attempts. Try again in 30 minutes.')
                    return self.render_to_response(self.get_context_data(form=form))

        return super().dispatch(request, *args, **kwargs)

    def form_invalid(self, form):
        username = self.request.POST.get('username', '')

        # only count attempts for real users
        if User.objects.filter(username=username).exists():
            attempts_key = f'login_attempts_{username}'
            lock_key = f'login_lock_{username}'

            # count failed attempt
            attempts = cache.get(attempts_key, 0) + 1
            cache.set(attempts_key, attempts, timeout=1800)

            remaining = 5 - attempts

            # lock after 5 failed attempts
            if attempts >= 5:
                cache.set(lock_key, True, timeout=1800)
                cache.delete(attempts_key)
                form.add_error(None, 'Too many failed login attempts. Try again in 30 minutes.')
            else:
                form.add_error(None, f'Invalid username or password. {remaining} attempts remaining.')

        else:
            # generic message for unknown usernames
            form.add_error(None, 'Invalid username or password.')

        return self.render_to_response(self.get_context_data(form=form))

    def form_valid(self, form):
        username = self.request.POST.get('username', '')

        # clear attempts after successful login
        cache.delete(f'login_attempts_{username}')
        cache.delete(f'login_lock_{username}')

        return super().form_valid(form)
    
@login_required
@never_cache
def teams(request):
    all_teams = Team.objects.select_related('department', 'teamLeader').all().prefetch_related('membership_set').order_by('teamName')
    active_teams = all_teams.filter(teamStatus='active')
    disbanded_teams = all_teams.filter(teamStatus='disbanded')
    return render(request, 'teams.html', {'teams': all_teams, 'active_teams': active_teams, 'disbanded_teams': disbanded_teams})

@login_required
@never_cache
def team_detail(request, team_id):
    team = get_object_or_404(Team, id=team_id)

    members = Membership.objects.filter(team=team)
    repos = Repository.objects.filter(team=team)
    contacts = ContactChannel.objects.filter(team=team)
    updep = Dependency.objects.filter(upstreamDep=team)
    downdep = Dependency.objects.filter(downstreamDep=team)
    return render(request, 'team_detail.html', {
        'team': team,
        'members': members,
        'repositories': repos,
        'contacts': contacts,
        'upstream_dependencies': updep,
        'downstream_dependencies': downdep
    })


@login_required
@never_cache
def organisation(request):
    departments = Department.objects.prefetch_related('team_set').filter(team__teamStatus='active').distinct().order_by('deptName')
    return render(request, 'organisation.html', {'departments': departments})




# Messaging 
@login_required
def conversations(request):
    delete_empty_conversations()
    conversations = Conversation.objects.filter(participants=request.user, messages__isnull=False).exclude(hidden_for=request.user).distinct().order_by('-lastUpdated')
    users = User.objects.exclude(id=request.user.id).order_by('username')
    hidden_for = models.ManyToManyField(User, related_name='hidden_conversations', blank=True)
    return render(request, 'messaging/conversations.html', {'conversations': conversations, 'users': users})

@login_required
def start_conversation(request, user_id):
    other_user = get_object_or_404(User, id=user_id)
    convos = Conversation.objects.filter(participants=request.user).filter(participants=other_user)
    convo = convos.first()
    # solves duplication glitch
    if convos.count() > 1:
        convos.exclude(id=convo.id).delete()
    if convo:
        convo.hidden_for.remove(request.user)
    else:
        convo = Conversation.objects.create()
        convo.participants.add(request.user, other_user)
    return redirect('chat', conversation_id=convo.id)

@login_required
def chat(request, conversation_id):
    convo = Conversation.objects.filter(id=conversation_id).first()
    if not convo:
        return redirect('conversations')
    if request.user not in convo.participants.all():
        return redirect('conversations')
    messages = Message.objects.filter(
        conversation=convo,
        is_draft=False
    ).order_by('timestamp')
    users = User.objects.exclude(id=request.user.id).order_by('username')
    conversations = Conversation.objects.filter(participants=request.user).exclude(hidden_for=request.user).distinct().order_by('-lastUpdated')
    if request.method == 'POST':
        content = request.POST.get('content')
        if content:
            is_draft = request.POST.get('is_draft') == 'true'

            if is_draft:
                Message.objects.create(
                    conversation=convo,
                    sender=request.user,
                    content=content,
                    is_draft=True
                )

            else:
                # delete the user's old draft for this chat
                Message.objects.filter(
                    conversation=convo,
                    sender=request.user,
                    is_draft=True
                ).delete()

                # create the actual sent message
                Message.objects.create(
                    conversation=convo,
                    sender=request.user,
                    content=content,
                    is_draft=False
                )

                # notify the other person
                for user in convo.participants.exclude(id=request.user.id):
                    Notification.objects.create(
                        user=user,
                        message=f"{request.user.username} sent you a message"
                    )
        return redirect('chat', conversation_id=convo.id)
    
    draft = Message.objects.filter(
        conversation=convo,
        sender=request.user,
        is_draft=True
    ).order_by('-timestamp').first()
    return render(request, 'messaging/chat.html', {'conversation': convo,'messages': messages,'users': users,'conversations': conversations,'draft': draft,})

def delete_empty_conversations():
    empty_convos = Conversation.objects.filter(messages__isnull=True)
    empty_convos.delete()

@login_required
def hide_conversation(request, conversation_id):
    convo = get_object_or_404(Conversation, id=conversation_id)
    convo.hidden_for.add(request.user)
    return redirect('inbox')


#excel report view
@login_required
def export_excel_report(request):
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    report_data = get_report_data()

    workbook = Workbook()


    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_sheet.append(["Category", "Count"])
    summary_sheet.append(["Users", report_data['users_count']])
    summary_sheet.append(["Teams", report_data['teams_count']])
    summary_sheet.append(["Departments / Organisations", report_data['departments_count']])
    summary_sheet.append(["Messages", report_data['messages_count']])
    summary_sheet.append(["Conversations", report_data['conversations_count']])


    chart = BarChart()
    chart.title = "System Overview"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Category"

    data = Reference(summary_sheet, min_col=2, min_row=1, max_row=6)
    categories = Reference(summary_sheet, min_col=1, min_row=2, max_row=6)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    summary_sheet.add_chart(chart, "D2")


    messages_sheet = workbook.create_sheet(title="Recent Chat History")
    messages_sheet.append(["Sender", "Conversation ID", "Message", "Timestamp"])

    for msg in report_data['recent_messages']:
        messages_sheet.append([
            msg.sender.username if msg.sender else "Unknown",
            msg.conversation.id if msg.conversation else "",
            msg.content,
            str(msg.timestamp),
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="system_report.xlsx"'

    workbook.save(response)
    return response

#pdf report view
@login_required
def export_pdf_report(request):
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    report_data = get_report_data()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="system_report.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    width, height = letter

    p.setFont("Helvetica-Bold", 18)
    p.drawString(50, height - 50, "System Report")

    p.setFont("Helvetica", 12)
    y = height - 90

    p.drawString(50, y, f"Users: {report_data['users_count']}")
    y -= 20
    p.drawString(50, y, f"Teams: {report_data['teams_count']}")
    y -= 20
    p.drawString(50, y, f"Departments / Organisations: {report_data['departments_count']}")
    y -= 20
    p.drawString(50, y, f"Messages: {report_data['messages_count']}")
    y -= 20
    p.drawString(50, y, f"Conversations: {report_data['conversations_count']}")
    y -= 40


    drawing = Drawing(400, 200)
    chart = VerticalBarChart()
    chart.x = 50
    chart.y = 50
    chart.height = 125
    chart.width = 300
    chart.data = [[
        report_data['users_count'],
        report_data['teams_count'],
        report_data['departments_count'],
        report_data['messages_count'],
        report_data['conversations_count'],
    ]]
    chart.categoryAxis.categoryNames = ['Users', 'Teams', 'Departments', 'Messages', 'Chats']
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = max(
        5,
        report_data['users_count'],
        report_data['teams_count'],
        report_data['departments_count'],
        report_data['messages_count'],
        report_data['conversations_count'],
    ) + 2
    chart.valueAxis.valueStep = 1

    drawing.add(chart)
    drawing.drawOn(p, 50, y - 170)

    y -= 210
    p.setFont("Helvetica-Bold", 14)
    p.drawString(50, y, "Recent Chat History")
    y -= 25

    p.setFont("Helvetica", 10)

    for msg in report_data['recent_messages']:
        line = f"{msg.sender.username}: {msg.content[:60]} ({msg.timestamp})"
        p.drawString(50, y, line[:100])
        y -= 15

        if y < 50:
            p.showPage()
            p.setFont("Helvetica", 10)
            y = height - 50

    p.showPage()
    p.save()
    return response
#report page view
@login_required
def reports_page(request):
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    report_data = get_report_data()

    chart_labels = ['Users', 'Teams', 'Departments', 'Messages', 'Conversations']
    chart_values = [
        report_data['users_count'],
        report_data['teams_count'],
        report_data['departments_count'],
        report_data['messages_count'],
        report_data['conversations_count'],
    ]

    return render(request, 'reports.html', {
        'report_data': report_data,
        'chart_labels': chart_labels,
        'chart_values': chart_values,
    })

    return render(request, 'reports.html', {
        'report_data': report_data,
        'chart_labels': chart_labels,
        'chart_values': chart_values,
    })
#full chat in excel
@login_required
def export_full_chat_excel(request):
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    grouped_chats = get_grouped_chat_data()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Full Chat Export"

    row_num = 1

    for chat in grouped_chats:
        sheet.cell(row=row_num, column=1, value=f"Chat: {chat['participant_names']}")
        row_num += 1

        sheet.cell(row=row_num, column=1, value="Sender")
        sheet.cell(row=row_num, column=2, value="Message")
        sheet.cell(row=row_num, column=3, value="Timestamp")
        row_num += 1

        for msg in chat['messages']:
            sheet.cell(row=row_num, column=1, value=msg.sender.username if msg.sender else "Unknown")
            sheet.cell(row=row_num, column=2, value=msg.content)
            sheet.cell(row=row_num, column=3, value=str(msg.timestamp))
            row_num += 1

        row_num += 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="grouped_chat_export.xlsx"'

    workbook.save(response)
    return response
#full chat in pdf
@login_required
def export_full_chat_pdf(request):
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    grouped_chats = get_grouped_chat_data()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="grouped_chat_export.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    width, height = letter
    y = height - 40

    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, y, "Full Chat Export")
    y -= 30

    for chat in grouped_chats:
        if y < 100:
            p.showPage()
            y = height - 40

        p.setFont("Helvetica-Bold", 12)
        p.drawString(50, y, f"Chat: {chat['participant_names']}")
        y -= 20

        p.setFont("Helvetica", 10)

        for msg in chat['messages']:
            sender = msg.sender.username if msg.sender else "Unknown"
            timestamp = str(msg.timestamp)
            content = msg.content if msg.content else ""

            header_line = f"{sender} - {timestamp}"
            p.drawString(60, y, header_line[:100])
            y -= 12

            chunk_size = 90
            for i in range(0, len(content), chunk_size):
                if y < 50:
                    p.showPage()
                    y = height - 40
                    p.setFont("Helvetica", 10)

                p.drawString(70, y, content[i:i + chunk_size])
                y -= 12

            y -= 10

            if y < 50:
                p.showPage()
                y = height - 40
                p.setFont("Helvetica", 10)

        y -= 10

    p.save()
    return response


@login_required
def drafts_page(request):
    drafts = Message.objects.filter(
        sender=request.user,
        is_draft=True
    ).order_by('-timestamp')

    return render(request, 'messaging/drafts.html', {'drafts': drafts})


@login_required
def edit_draft(request, draft_id):
    draft = get_object_or_404(
        Message,
        id=draft_id,
        sender=request.user,
        is_draft=True
    )

    return redirect('chat', conversation_id=draft.conversation.id)


@login_required
def delete_draft(request, draft_id):
    draft = get_object_or_404(
        Message,
        id=draft_id,
        sender=request.user,
        is_draft=True
    )

    if request.method == "POST":
        draft.delete()

    return redirect('drafts_page')

@login_required
def mark_notifications_read(request):
    request.user.notification_set.filter(is_read=False).update(is_read=True)
    return JsonResponse({'status': 'ok'})

from django.http import JsonResponse

@login_required
def mark_notifications_read(request):
    request.user.notification_set.filter(is_read=False).update(is_read=True)
    return JsonResponse({'status': 'ok'})

@login_required
def updateprofile(request):
    if request.method == 'POST':
        user_form = UserForm(request.POST, instance=request.user)
        if user_form.is_valid():
            user_form.save()
            return redirect('updateprofile')
    else:
        user_form = UserForm(instance=request.user)
    return render(request, 'accounts/profile.html', {'user_form': user_form})
