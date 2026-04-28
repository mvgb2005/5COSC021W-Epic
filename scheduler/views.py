from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.views.decorators.cache import never_cache
from django.contrib import messages
from django.db.models import Q
from .models import ScheduleMeeting
from .forms import ScheduleMeetingForm
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


@login_required
@never_cache
def schedule_page(request):
    # get search and filter values
    q = request.GET.get('q') if request.GET.get('q') != None else ''
    schedule_filter = request.GET.get('filter') if request.GET.get('filter') != None else ''

    # only show meetings created by the logged-in user
    meetings = ScheduleMeeting.objects.filter(createdBy=request.user)

    # search title and message using icontains
    meetings = meetings.filter(
        Q(meetingTitle__icontains=q) |
        Q(meetingMessage__icontains=q)
    ).order_by('meetingDate', 'meetingTime')

    # filter by weekly, monthly, or upcoming
    if schedule_filter:
        meetings = meetings.filter(scheduleType=schedule_filter)

    # split meetings into sections
    upcoming_meetings = meetings.filter(scheduleType='upcoming')
    weekly_meetings = meetings.filter(scheduleType='weekly')
    monthly_meetings = meetings.filter(scheduleType='monthly')

    # send data to template
    context = {
        'meetings': meetings,
        'upcoming_meetings': upcoming_meetings,
        'weekly_meetings': weekly_meetings,
        'monthly_meetings': monthly_meetings,
        'q': q,
        'schedule_filter': schedule_filter,
    }

    return render(request, 'scheduler/schedule.html', context)


@login_required
@never_cache
def create_schedule(request):
    # create new meeting
    if request.method == 'POST':
        form = ScheduleMeetingForm(request.POST)

        if form.is_valid():
            meeting = form.save(commit=False)
            meeting.createdBy = request.user
            meeting.save()
            messages.success(request, 'Meeting added successfully.')
            return redirect('schedule')
    else:
        form = ScheduleMeetingForm()

    return render(request, 'scheduler/schedule_form.html', {'form': form})


@login_required
@never_cache
def update_schedule(request, schedule_id):
    # only allow owner to update their meeting
    meeting = get_object_or_404(ScheduleMeeting, id=schedule_id, createdBy=request.user)

    if request.method == 'POST':
        form = ScheduleMeetingForm(request.POST, instance=meeting)

        if form.is_valid():
            form.save()
            messages.success(request, 'Meeting updated successfully.')
            return redirect('schedule')
    else:
        form = ScheduleMeetingForm(instance=meeting)

    return render(request, 'scheduler/schedule_form.html', {'form': form, 'meeting': meeting})


@login_required
@never_cache
def delete_schedule(request, schedule_id):
    # only allow owner to delete their meeting
    meeting = get_object_or_404(ScheduleMeeting, id=schedule_id, createdBy=request.user)

    if request.method == 'POST':
        meeting.delete()
        messages.success(request, 'Meeting deleted successfully.')
        return redirect('schedule')

    return render(request, 'scheduler/schedule_delete.html', {'meeting': meeting})


# filter schedule report data
def get_filtered_schedule_meetings(request):
    # get report filters
    q = request.GET.get('q') if request.GET.get('q') != None else ''
    schedule_filter = request.GET.get('filter') if request.GET.get('filter') != None else ''
    platform_filter = request.GET.get('platform') if request.GET.get('platform') != None else ''
    date_filter = request.GET.get('date') if request.GET.get('date') != None else ''

    # get all meetings for admin report
    meetings = ScheduleMeeting.objects.all().order_by('meetingDate', 'meetingTime')

    # search title and message
    meetings = meetings.filter(
        Q(meetingTitle__icontains=q) |
        Q(meetingMessage__icontains=q)
    )

    # filter by schedule type
    if schedule_filter:
        meetings = meetings.filter(scheduleType=schedule_filter)

    # filter by platform
    if platform_filter:
        meetings = meetings.filter(meetingPlatform=platform_filter)

    # filter by date
    if date_filter:
        meetings = meetings.filter(meetingDate=date_filter)

    return meetings, q, schedule_filter, platform_filter, date_filter


@login_required
def schedule_reports(request):
    # only admin users can view schedule reports
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    # get filtered meetings
    meetings, q, schedule_filter, platform_filter, date_filter = get_filtered_schedule_meetings(request)

    # count meeting types
    upcoming_count = meetings.filter(scheduleType='upcoming').count()
    weekly_count = meetings.filter(scheduleType='weekly').count()
    monthly_count = meetings.filter(scheduleType='monthly').count()

    # count platforms
    zoom_count = meetings.filter(meetingPlatform='zoom').count()
    teams_count = meetings.filter(meetingPlatform='teams').count()
    google_count = meetings.filter(meetingPlatform='google_meet').count()

    # send report data to template
    context = {
        'meetings': meetings,
        'q': q,
        'schedule_filter': schedule_filter,
        'platform_filter': platform_filter,
        'date_filter': date_filter,
        'upcoming_count': upcoming_count,
        'weekly_count': weekly_count,
        'monthly_count': monthly_count,
        'zoom_count': zoom_count,
        'teams_count': teams_count,
        'google_count': google_count,
    }

    return render(request, 'scheduler/schedule_reports.html', context)


@login_required
def export_schedule_excel(request):
    # only admin users can download excel
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    # export only filtered meetings
    meetings, q, schedule_filter, platform_filter, date_filter = get_filtered_schedule_meetings(request)

    # create workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Schedule Report"

    # add headings
    sheet.append(["Title", "Date", "Time", "Platform", "Type", "Message", "Created By"])

    # add meeting rows
    for meeting in meetings:
        sheet.append([
            meeting.meetingTitle,
            str(meeting.meetingDate),
            str(meeting.meetingTime),
            meeting.get_meetingPlatform_display(),
            meeting.get_scheduleType_display(),
            meeting.meetingMessage,
            meeting.createdBy.username,
        ])

    # create summary sheet
    summary = workbook.create_sheet(title="Summary")
    summary.append(["Category", "Count"])
    summary.append(["Upcoming", meetings.filter(scheduleType='upcoming').count()])
    summary.append(["Weekly", meetings.filter(scheduleType='weekly').count()])
    summary.append(["Monthly", meetings.filter(scheduleType='monthly').count()])

    # create bar chart
    chart = BarChart()
    chart.title = "Schedule Overview"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Type"

    data = Reference(summary, min_col=2, min_row=1, max_row=4)
    cats = Reference(summary, min_col=1, min_row=2, max_row=4)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    summary.add_chart(chart, "D2")

    # return excel file
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="schedule_report.xlsx"'

    workbook.save(response)
    return response


@login_required
def export_schedule_pdf(request):
    # only admin users can download pdf
    if not request.user.is_staff:
        return render(request, 'admin_only.html')

    # export only filtered meetings
    meetings, q, schedule_filter, platform_filter, date_filter = get_filtered_schedule_meetings(request)

    # create pdf response
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="schedule_report.pdf"'

    p = canvas.Canvas(response, pagesize=letter)
    width, height = letter
    y = height - 50

    # pdf title
    p.setFont("Helvetica-Bold", 18)
    p.drawString(50, y, "Schedule Report")
    y -= 35

    p.setFont("Helvetica", 11)

    # add meetings to pdf
    for meeting in meetings:
        if y < 80:
            p.showPage()
            y = height - 50
            p.setFont("Helvetica", 11)

        p.drawString(50, y, f"Title: {meeting.meetingTitle}")
        y -= 15
        p.drawString(50, y, f"Date: {meeting.meetingDate}  Time: {meeting.meetingTime}")
        y -= 15
        p.drawString(50, y, f"Platform: {meeting.get_meetingPlatform_display()}  Type: {meeting.get_scheduleType_display()}")
        y -= 15
        p.drawString(50, y, f"Created By: {meeting.createdBy.username}")
        y -= 15
        p.drawString(50, y, f"Message: {meeting.meetingMessage[:90]}")
        y -= 25

    # save pdf
    p.save()
    return response